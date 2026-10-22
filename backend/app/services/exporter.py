"""导出 xlsx、csv（04 第 9 节）。

按 seq 输出范围内所有条目（包括失败、跳过和未处理的）；先原始列、后结果列；结果列与原始列重名时加“（文枢）”。
分页读取（每次 2,000 行），写文件在工作线程里做，不阻塞事件循环；文件先写到 data/exports/ 再返回。

单元格清理（两种格式都做，防止导出失败或被 Excel 当成公式执行）：去掉 Excel 不允许的控制字符；
xlsx 除数值列（置信度、字数、压缩比）外一律按文本写入；csv 里像公式的值前加英文单引号；超过 32,767 字截断。
"""

from __future__ import annotations

import asyncio
import csv
import re
import uuid
from collections.abc import Callable
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from app.context import AppContext
from app.db.models import JobItemRow, JobRow
from app.db.repo import items as items_repo
from app.schemas.common import TASK_TYPE_LABELS
from app.services import review
from app.services.job_service import aliases_of, require_job
from app.utils import loads, local_now, parse_iso

EXPORT_PAGE = 2000
# 与 openpyxl 的 ILLEGAL_CHARACTERS_RE 相同：\x00–\x08、\x0b、\x0c、\x0e–\x1f，写进 xlsx 会抛 IllegalCharacterError
ILLEGAL_CHARACTERS_RE = re.compile(r"[\000-\010]|[\013-\014]|[\016-\037]")
MAX_CELL_CHARS = 32767
TRUNCATED_SUFFIX = "……（已截断）"
CSV_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
_PURE_NUMBER = re.compile(r"^[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?$")
STATUS_TEXT = {"succeeded": "成功", "failed": "失败", "skipped": "跳过", "pending": "未处理", "running": "未处理"}
_WINDOWS_INVALID = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
_CJK_WIDE = re.compile(r"[ᄀ-ᅟ⺀-꓏가-힣豈-﫿︰-﹏＀-｠￠-￦]")
MEDIA_TYPES = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv; charset=utf-8",
}


@dataclass
class ExportColumn:
    header: str
    wrap: bool = False
    number_format: str | None = None
    csv_format: Callable[[Any], str] | None = None
    numeric: bool = False          # 数值列按数字写入，其余一律按文本


def clean_cell_text(value: str) -> str:
    """去掉 Excel 不允许的控制字符；超过 32,767 字时截断并在末尾加“……（已截断）”。"""
    text = ILLEGAL_CHARACTERS_RE.sub("", value)
    if len(text) > MAX_CELL_CHARS:
        text = text[: MAX_CELL_CHARS - len(TRUNCATED_SUFFIX)] + TRUNCATED_SUFFIX
    return text


def csv_cell_text(value: str) -> str:
    """csv：值以 =、+、-、@、制表符或回车开头且不是纯数字时，在前面加一个英文单引号。"""
    text = clean_cell_text(value)
    if text.startswith(CSV_FORMULA_PREFIXES) and not _PURE_NUMBER.match(text):
        text = "'" + text
        if len(text) > MAX_CELL_CHARS:
            text = text[: MAX_CELL_CHARS - len(TRUNCATED_SUFFIX)] + TRUNCATED_SUFFIX
    return text


def _fmt2(value: Any) -> str:
    return f"{value:.2f}" if isinstance(value, int | float) else ""


def _fmt1(value: Any) -> str:
    return f"{value:.1f}" if isinstance(value, int | float) else ""


def safe_filename(name: str) -> str:
    cleaned = _WINDOWS_INVALID.sub("", name).strip().rstrip(". ")
    return cleaned or "导出"


def _local_time(value: str | None) -> str:
    if not value:
        return ""
    try:
        return parse_iso(value).astimezone().strftime("%Y-%m-%d %H:%M:%S")
    except ValueError:
        return value


def _display_width(text: str) -> int:
    first_line = text.split("\n", 1)[0]
    return len(first_line) + len(_CJK_WIDE.findall(first_line))


def _is_lines(source: dict[str, Any]) -> bool:
    columns = source.get("columns") or []
    return len(columns) == 1 and columns[0].get("key") == "text" and columns[0].get("letter") is None


def build_columns(job: JobRow) -> tuple[list[ExportColumn], bool]:
    source = loads(job.source_json, {}) or {}
    lines = _is_lines(source)
    if lines:
        originals = [ExportColumn("行号"), ExportColumn("文本", wrap=True)]
    else:
        originals = [ExportColumn(str(c.get("title") or c.get("key") or "")) for c in source.get("columns") or []]
        text_keys = set(source.get("text_columns") or [])
        for col, spec in zip(originals, source.get("columns") or [], strict=False):
            col.wrap = spec.get("key") in text_keys
    if job.task_type == "classify":
        results = [
            ExportColumn("意图"),
            ExportColumn("修正前意图"),
            ExportColumn("置信度", number_format="0.00", csv_format=_fmt2, numeric=True),
            ExportColumn("判断依据", wrap=True),
            ExportColumn("人工修正"),
            ExportColumn("已复核"),
            ExportColumn("处理状态"),
            ExportColumn("错误信息", wrap=True),
        ]
    else:
        results = [
            ExportColumn("摘要", wrap=True),
            ExportColumn("修正前摘要", wrap=True),
            ExportColumn("关键要点", wrap=True),
            ExportColumn("原文字数", numeric=True),
            ExportColumn("摘要字数", numeric=True),
            ExportColumn("压缩比", number_format="0.0", csv_format=_fmt1, numeric=True),
            ExportColumn("人工修正"),
            ExportColumn("已复核"),
            ExportColumn("处理状态"),
            ExportColumn("错误信息", wrap=True),
        ]
    taken = {c.header for c in originals}
    for col in results:
        if col.header in taken:
            col.header = f"{col.header}（文枢）"
    return originals + results, lines


def _key_points_text(value: str | None) -> str:
    points = loads(value, []) or []
    lines: list[str] = []
    for index, point in enumerate(points, start=1):
        if not isinstance(point, dict) or not point.get("text"):
            continue
        sources = [str(s) for s in point.get("sources") or []]
        suffix = f"（第 {'、'.join(sources)} 段）" if sources else ""
        lines.append(f"{index}. {point['text']}{suffix}")
    return "\n".join(lines)


def row_values(job: JobRow, item: JobItemRow, aliases: dict[str, str], lines: bool) -> list[Any]:
    source = loads(job.source_json, {}) or {}
    values = loads(item.row_values_json, {}) or {}
    if lines:
        original: list[Any] = [item.row_number, values.get("text", item.input_text)]
    else:
        original = [values.get(str(c.get("key")), "") for c in source.get("columns") or []]
    status = STATUS_TEXT.get(item.status, item.status)
    error = (item.error_message or "") if item.status == "failed" else ""
    if job.task_type == "classify":
        label = item.label
        model_label = (aliases.get(label, label) if label else "") or ""
        effective = item.manual_label or model_label
        succeeded = item.status == "succeeded"
        result = [
            effective or "",
            model_label if item.manual_label else "",          # 修正前：模型给出的类别
            item.confidence if succeeded and item.confidence is not None else None,
            (item.reason or "") if succeeded else "",
            "是" if item.manual_label else "",
            "是" if item.reviewed else "",
            status,
            error,
        ]
    else:
        succeeded = item.status == "succeeded"
        edited = item.manual_summary is not None
        summary_text = item.manual_summary if edited else ((item.summary or "") if succeeded else "")
        output_chars = review.effective_output_chars(item.output_chars, item.summary, item.manual_summary)
        ratio = round(item.input_chars / output_chars, 1) if succeeded and output_chars else None
        result = [
            summary_text or "",
            (item.summary or "") if edited else "",            # 修正前：模型给出的摘要
            _key_points_text(item.key_points_json) if succeeded else "",
            item.input_chars,
            output_chars if succeeded else None,
            ratio,
            "是" if edited else "",
            "是" if item.reviewed else "",
            status,
            error,
        ]
    return original + result


class _XlsxExport:
    def __init__(self, path: Path, columns: list[ExportColumn]) -> None:
        from openpyxl import Workbook

        self.path = path
        self.columns = columns
        self.wb = Workbook(write_only=True)
        self.ws = self.wb.create_sheet("结果")
        self.started = False

    def _start(self, sample: list[list[Any]]) -> None:
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        self._wrap = Alignment(wrap_text=True, vertical="top")
        for index, col in enumerate(self.columns):
            widths = [_display_width(col.header)]
            widths.extend(_display_width("" if row[index] is None else str(row[index])) for row in sample[:2000])
            width = min(60, max(10, max(widths) + 2))
            self.ws.column_dimensions[get_column_letter(index + 1)].width = width
        try:
            self.ws.freeze_panes = "A2"
        except (AttributeError, TypeError, ValueError):
            pass
        bold = Font(bold=True)
        fill = PatternFill("solid", fgColor="EEF2F6")
        header: list[Any] = []
        for col in self.columns:
            cell = WriteOnlyCell(self.ws, value=clean_cell_text(col.header))
            cell.data_type = "s"
            cell.font = bold
            cell.fill = fill
            header.append(cell)
        self.ws.append(header)
        self.started = True

    def _text_cell(self, value: Any, *, wrap: bool) -> Any:
        """按文本写入：以 = 开头的文本、需要换行的列用 data_type = "s" 的单元格，其余字符串本来就是文本。"""
        from openpyxl.cell import WriteOnlyCell

        text = clean_cell_text(str(value))
        if not wrap and not text.startswith("="):
            return text
        cell = WriteOnlyCell(self.ws, value=text)
        cell.data_type = "s"
        if wrap:
            cell.alignment = self._wrap
        return cell

    def write_rows(self, rows: list[list[Any]]) -> None:
        from openpyxl.cell import WriteOnlyCell

        if not self.started:
            self._start(rows)
        for values in rows:
            out: list[Any] = []
            for col, value in zip(self.columns, values, strict=True):
                if value is None:
                    out.append(None)
                elif col.numeric and isinstance(value, int | float):
                    if col.number_format:
                        cell = WriteOnlyCell(self.ws, value=value)
                        cell.number_format = col.number_format
                        out.append(cell)
                    else:
                        out.append(value)
                else:
                    out.append(self._text_cell(value, wrap=col.wrap))
            self.ws.append(out)

    def close(self, info: list[tuple[str, Any]]) -> None:
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Font, PatternFill

        if not self.started:
            self._start([])
        sheet = self.wb.create_sheet("任务信息")
        sheet.column_dimensions["A"].width = 16
        sheet.column_dimensions["B"].width = 60
        bold = Font(bold=True)
        fill = PatternFill("solid", fgColor="EEF2F6")
        header = []
        for title in ("项目", "内容"):
            cell = WriteOnlyCell(sheet, value=title)
            cell.font = bold
            cell.fill = fill
            header.append(cell)
        sheet.append(header)
        for key, value in info:
            row: list[Any] = [key]
            if isinstance(value, str):
                cell = WriteOnlyCell(sheet, value=clean_cell_text(value))
                cell.data_type = "s"
                row.append(cell)
            else:
                row.append(value)
            sheet.append(row)
        self.wb.save(str(self.path))


class _CsvExport:
    """UTF-8 带 BOM，逗号分隔，CRLF 换行，按需加引号。"""

    def __init__(self, path: Path, columns: list[ExportColumn]) -> None:
        self.columns = columns
        self.handle = path.open("w", encoding="utf-8-sig", newline="")
        self.writer = csv.writer(self.handle, lineterminator="\r\n")
        self.writer.writerow([csv_cell_text(c.header) for c in columns])

    def write_rows(self, rows: list[list[Any]]) -> None:
        for values in rows:
            out: list[str] = []
            for col, value in zip(self.columns, values, strict=True):
                if value is None:
                    out.append("")
                elif col.csv_format is not None:
                    out.append(col.csv_format(value))
                elif col.numeric and isinstance(value, int | float):
                    out.append(str(value))
                else:
                    out.append(csv_cell_text(str(value)))
            self.writer.writerow(out)

    def close(self, _info: list[tuple[str, Any]]) -> None:
        self.handle.close()


def _info_rows(job: JobRow, counts: dict[str, int], scope: str) -> list[tuple[str, Any]]:
    model = loads(job.model_snapshot_json, {}) or {}
    prompt = loads(job.prompt_snapshot_json, {}) or {}
    pending = counts.get("pending", 0) + counts.get("running", 0)
    total_tokens = int(job.prompt_tokens) + int(job.completion_tokens)
    return [
        ("任务名", job.name),
        ("文件名", job.file_name),
        ("任务类型", TASK_TYPE_LABELS.get(job.task_type, job.task_type)),
        ("模型", f"{model.get('name', '')}（{model.get('model', '')}）"),
        ("提示词", f"{prompt.get('name', '')} v{job.prompt_version}"),
        ("并发", int(job.concurrency)),
        ("开始时间", _local_time(job.started_at)),
        ("结束时间", _local_time(job.finished_at) or "未结束"),
        ("成功", counts.get("succeeded", 0)),
        ("失败", counts.get("failed", 0)),
        ("跳过", counts.get("skipped", 0)),
        ("未处理", pending),
        ("tokens 合计", f"{total_tokens:,}（输入 {int(job.prompt_tokens):,}，输出 {int(job.completion_tokens):,}）"),
        ("导出范围", "全部条目" if scope == "all" else "当前筛选结果"),
        ("导出时间", local_now().strftime("%Y-%m-%d %H:%M:%S")),
    ]


def download_name(job: JobRow, fmt: str, when: datetime | None = None) -> str:
    stem = Path(job.file_name).stem or job.file_name
    stamp = (when or local_now()).strftime("%Y%m%d-%H%M")
    return safe_filename(f"{stem}_{TASK_TYPE_LABELS.get(job.task_type, job.task_type)}_{stamp}") + f".{fmt}"


async def export_job(
    ctx: AppContext, job_id: int, fmt: str, filters: items_repo.ItemFilters, scope: str
) -> tuple[Path, str]:
    async with ctx.db.session() as session:
        job = await require_job(session, job_id)
        counts = await items_repo.status_counts(session, job_id)
    aliases = aliases_of(job)
    columns, lines = build_columns(job)
    ctx.cfg.exports_dir.mkdir(parents=True, exist_ok=True)
    path = ctx.cfg.exports_dir / f"job{job_id}-{uuid.uuid4().hex[:10]}.{fmt}"
    output: _XlsxExport | _CsvExport = (
        await asyncio.to_thread(_XlsxExport, path, columns) if fmt == "xlsx" else await asyncio.to_thread(_CsvExport, path, columns)
    )
    try:
        after = 0
        while True:
            async with ctx.db.session() as session:
                page = await items_repo.export_page(session, job_id, filters, aliases, after, EXPORT_PAGE)
            if not page:
                break
            rows = [row_values(job, item, aliases, lines) for item in page]
            await asyncio.to_thread(output.write_rows, rows)
            after = page[-1].seq
        await asyncio.to_thread(output.close, _info_rows(job, counts, scope))
    except BaseException:
        if isinstance(output, _CsvExport):
            output.handle.close()
        path.unlink(missing_ok=True)
        raise
    return path, download_name(job, fmt)

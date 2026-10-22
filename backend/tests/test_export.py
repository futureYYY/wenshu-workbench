"""导出 xlsx、csv（04 第 9 节；08 第 10.1 节“导出”）。"""

from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any
from urllib.parse import unquote

from openpyxl import load_workbook

from .conftest import Api

CLASSIFY_RESULT = ["意图", "修正前意图", "置信度", "判断依据", "人工修正", "已复核", "处理状态", "错误信息"]
SUMMARIZE_RESULT = [
    "摘要", "修正前摘要", "关键要点", "原文字数", "摘要字数", "压缩比", "人工修正", "已复核", "处理状态", "错误信息",
]


def _download(api: Api, job_id: int, **params: Any) -> tuple[bytes, str]:
    resp = api.get(f"/api/jobs/{job_id}/export", params=params)
    assert resp.status_code == 200, resp.text
    disposition = resp.headers["content-disposition"]
    assert disposition.startswith("attachment; filename*=UTF-8''")
    return resp.content, unquote(disposition.split("''", 1)[1])


def test_classify_xlsx(api: Api, fast_model: dict[str, Any], fixture_files: dict[str, Path]) -> None:
    file = api.upload(fixture_files["intent_200.xlsx"])
    job = api.create_job(file, model_id=fast_model["id"], sheet="对话", text_columns=["用户问题"], limit=30)
    done = api.wait_job(job["id"])
    content, name = _download(api, job["id"], format="xlsx")
    assert name.startswith("intent_200_意图识别_") and name.endswith(".xlsx")

    wb = load_workbook(io.BytesIO(content))
    assert wb.sheetnames == ["结果", "任务信息"]
    sheet = wb["结果"]
    rows = list(sheet.iter_rows(values_only=True))
    assert list(rows[0]) == ["会话ID", "用户问题", "时间", "渠道", *CLASSIFY_RESULT]
    assert len(rows) - 1 == done["progress"]["total"]
    assert sheet.freeze_panes == "A2"
    assert sheet.cell(row=1, column=1).font.bold
    status_col = rows[0].index("处理状态")
    assert {r[status_col] for r in rows[1:]} <= {"成功", "跳过"}
    confidence_col = rows[0].index("置信度") + 1
    first_success = next(i for i, r in enumerate(rows[1:], start=2) if r[status_col] == "成功")
    assert sheet.cell(row=first_success, column=confidence_col).number_format == "0.00"
    info = {r[0]: r[1] for r in wb["任务信息"].iter_rows(min_row=2, values_only=True)}
    assert info["任务类型"] == "意图识别" and info["文件名"] == "intent_200.xlsx" and "tokens 合计" in info


def test_lines_csv_bom_crlf(api: Api, fast_model: dict[str, Any]) -> None:
    job = api.create_job(api.paste("导出甲\n\n导出乙"), model_id=fast_model["id"])
    api.wait_job(job["id"])
    content, name = _download(api, job["id"], format="csv")
    assert name.endswith(".csv")
    assert content.startswith(b"\xef\xbb\xbf")
    assert b"\r\n" in content
    rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
    assert rows[0] == ["行号", "文本", *CLASSIFY_RESULT]
    assert [r[0] for r in rows[1:]] == ["1", "2", "3"]
    assert rows[2][rows[0].index("处理状态")] == "跳过"


def test_result_column_name_collision(api: Api, fast_model: dict[str, Any], tmp_path: Path) -> None:
    path = tmp_path / "collision.csv"
    path.write_text("意图,内容\n原有的值,请问怎么退货\n", encoding="utf-8", newline="")
    file = api.upload(path)
    job = api.create_job(file, model_id=fast_model["id"], text_columns=["内容"])
    api.wait_job(job["id"])
    content, _ = _download(api, job["id"], format="csv")
    header = next(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
    assert header[:3] == ["意图", "内容", "意图（文枢）"]


def test_summarize_columns_and_key_points(api: Api, fast_model: dict[str, Any]) -> None:
    text = "第一段讲了宽带掉线的问题，用户很着急。\n第二段说明工程师上门检查。\n第三段写了问题已经解决。"
    job = api.create_job(api.paste(text), task_type="summarize", model_id=fast_model["id"])
    api.wait_job(job["id"])
    content, name = _download(api, job["id"], format="csv")
    assert "_文本摘要_" in name
    rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
    assert rows[0] == ["行号", "文本", *SUMMARIZE_RESULT]
    record = dict(zip(rows[0], rows[1], strict=True))
    assert record["摘要"].startswith("演示摘要：")
    assert record["关键要点"].splitlines()[0].startswith("1. ")
    assert "（第 " in record["关键要点"] and record["处理状态"] == "成功"
    assert float(record["压缩比"]) > 0


def test_filtered_scope(api: Api) -> None:
    model = api.mock_model("导出-坏 JSON", "bad_json,fast")
    job = api.create_job(api.paste("筛选导出甲\n\n筛选导出乙"), model_id=model["id"])
    api.wait_job(job["id"])
    content, _ = _download(api, job["id"], format="csv", scope="filtered", status="failed")
    rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
    assert len(rows) == 3 and all(r[rows[0].index("处理状态")] == "失败" for r in rows[1:])
    assert all(r[rows[0].index("错误信息")].startswith("输出无法解析") for r in rows[1:])
    everything, _ = _download(api, job["id"], format="csv", scope="all", status="failed")
    assert len(list(csv.reader(io.StringIO(everything.decode("utf-8-sig"))))) == 4


def test_cell_sanitizing(api: Api, fast_model: dict[str, Any], fixture_files: dict[str, Path]) -> None:
    """04 第 9 节单元格清理：以 =、+、-、@ 开头的文本，控制字符，超过 32,767 字的单元格（两种格式都要能导出）。"""
    file = api.upload(fixture_files["export_edge.csv"])
    job = api.create_job(file, model_id=fast_model["id"], text_columns=["留言"], limit=None)
    api.wait_job(job["id"], timeout=120)

    content, _ = _download(api, job["id"], format="xlsx", scope="all")
    sheet = load_workbook(io.BytesIO(content))["结果"]
    header = [c.value for c in sheet[1]]
    column = header.index("留言") + 1
    cells = {sheet.cell(row=r, column=1).value: sheet.cell(row=r, column=column) for r in range(2, sheet.max_row + 1)}
    formula = cells["1"]
    assert formula.value == "=SUM(A1)" and formula.data_type == "s"   # 按文本写入，不是公式
    assert cells["9"].data_type == "s" and cells["9"].value.startswith("=A1+B1")
    assert cells["6"].value == "账单里有个看不见的字符请帮我查一下"      # 去掉了 \x0b
    assert cells["7"].value == "退款进度什么时候到账"                    # 去掉了 \x1f
    long_cell = cells["8"].value
    assert len(long_cell) == 32767 and long_cell.endswith("……（已截断）")
    confidence = sheet.cell(row=2, column=header.index("置信度") + 1)
    assert confidence.data_type == "n"                                  # 数值列仍是数字

    content, _ = _download(api, job["id"], format="csv", scope="all")
    rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
    index = rows[0].index("留言")
    values = {row[0]: row[index] for row in rows[1:]}
    assert values["1"] == "'=SUM(A1)"
    assert values["2"].startswith("'+1 ")
    assert values["3"] == "-2"                                          # 纯数字不加单引号
    assert values["4"].startswith("'-2 号线")
    assert values["5"].startswith("'@x")
    assert values["6"] == "账单里有个看不见的字符请帮我查一下"
    assert len(values["8"]) == 32767 and values["8"].endswith("……（已截断）")
    assert values["10"] == "普通的一行：请问宽带怎么续费"


def test_clean_cell_text_unit() -> None:
    from app.services.exporter import clean_cell_text, csv_cell_text

    assert clean_cell_text("a\x00b\x08c\x0bd\x0ce\x0ef\x1fg\th\ni\rj") == "abcdefg\th\ni\rj"
    assert csv_cell_text("\t制表符开头") == "'\t制表符开头"
    assert csv_cell_text("+1.5e3") == "+1.5e3" and csv_cell_text("-0.25") == "-0.25"
    assert csv_cell_text("普通") == "普通"

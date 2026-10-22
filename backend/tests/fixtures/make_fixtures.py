"""生成测试数据（08 第 10.2 节）。大文件不提交，测试运行前自动生成。

单独运行：backend\\.venv\\Scripts\\python.exe backend\\tests\\fixtures\\make_fixtures.py
输出到 backend/tests/fixtures/generated/（已在 .gitignore 中忽略），可以拖进界面手动验收。
"""

from __future__ import annotations

import csv
import random
import sys
from datetime import datetime, timedelta
from pathlib import Path

DEFAULT_OUT = Path(__file__).resolve().parent / "generated"

INTENT_SAMPLES = {
    "咨询": [
        "请问流量套餐怎么收费",
        "会员积分什么时候过期",
        "退货运费谁来承担",
        "发票可以开电子版吗",
        "周末客服几点下班",
        "国际漫游需要提前开通吗",
        "宽带安装要多久",
    ],
    "办理": [
        "帮我开通国际漫游",
        "我要退订每月的视频会员",
        "把套餐改成 99 元的那档",
        "申请修改收货地址",
        "我想注销这个账号",
        "帮我补办一张电话卡",
    ],
    "投诉": [
        "快递三天没动了，客服也不回消息",
        "你们乱扣我的话费，我要投诉",
        "收到的商品是坏的，要求赔偿",
        "装宽带的师傅迟到两个小时，态度还很差",
        "退款申请了一周还没到账",
    ],
    "建议": [
        "建议 App 增加夜间模式",
        "希望能支持按天购买流量包",
        "客服排队太久，建议加一个回拨功能",
        "账单页面能不能显示每天的用量",
    ],
    "其他": ["你好", "嗯嗯", "谢谢", "123456", "哈哈哈"],
}

CHANNELS = ["App", "电话", "网页", "微信", "门店"]


def _questions(count: int, rng: random.Random) -> list[str]:
    """覆盖 5 类意图，约 10% 重复，3 行为空。"""
    pool = [(intent, text) for intent, texts in INTENT_SAMPLES.items() for text in texts]
    questions: list[str] = []
    for index in range(count):
        intent, text = pool[index % len(pool)]
        questions.append(f"{text}（第 {index + 1} 条）" if index >= len(pool) else text)
    for index in rng.sample(range(10, count), k=max(1, count // 10)):
        questions[index] = questions[index - rng.randint(1, 9)]
    for index in (count // 4, count // 2, count - 5):
        questions[index] = ""
    return questions


def make_intent_xlsx(path: Path, rows: int = 200) -> None:
    from openpyxl import Workbook

    rng = random.Random(20261021)
    wb = Workbook()
    ws = wb.active
    ws.title = "对话"
    ws.append(["会话ID", "用户问题", "时间", "渠道"])
    start = datetime(2026, 9, 15, 9, 0, 0)
    for index, question in enumerate(_questions(rows, rng)):
        ws.append([10023 + index, question, start + timedelta(minutes=7 * index), CHANNELS[index % len(CHANNELS)]])
    notes = wb.create_sheet("说明")
    notes.append(["项目", "内容"])
    notes.append(["来源", "测试数据"])
    notes.append(["行数", rows])
    notes.append(["意图", "咨询、办理、投诉、建议、其他"])
    notes.append(["备注", "约 10% 重复，3 行为空"])
    wb.save(path)


def make_gbk_csv(path: Path, rows: int = 1000) -> None:
    rng = random.Random(1)
    questions = _questions(rows, rng)
    with path.open("w", encoding="gb18030", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(["编号", "客户留言", "城市"])
        for index, question in enumerate(questions):
            writer.writerow([index + 1, question or "", rng.choice(["北京", "上海", "广州", "成都", "武汉"])])


def make_semicolon_csv(path: Path) -> None:
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh, delimiter=";")
        writer.writerow(["工单号", "问题描述", "优先级"])
        writer.writerow(["A001", "第一行\n第二行：引号内带换行", "高"])
        writer.writerow(["A002", "普通的一行；含中文分号", "中"])
        writer.writerow(["A003", "包含;英文分号的\"引号\"字段", "低"])
        writer.writerow(["A004", "", "低"])


def make_text_files(utf16: Path, bom: Path) -> None:
    lines = ["第一行：咨询流量", "", "第三行：办理退订", "第四行：投诉扣费"]
    utf16.write_text("\r\n".join(lines) + "\r\n", encoding="utf-16")
    bom.write_text("\n".join(lines) + "\n", encoding="utf-8-sig")


def make_legacy_xls(path: Path) -> None:
    import xlwt

    book = xlwt.Workbook(encoding="utf-8")
    sheet = book.add_sheet("旧版")
    for col, title in enumerate(["编号", "内容", "金额"]):
        sheet.write(0, col, title)
    for row in range(1, 11):
        sheet.write(row, 0, row)
        sheet.write(row, 1, f"旧版 Excel 第 {row} 行：请问账单怎么查")
        sheet.write(row, 2, row * 10.5)
    book.save(str(path))


def make_big_csv(path: Path, rows: int = 120_000) -> None:
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(["id", "text"])
        for index in range(rows):
            writer.writerow([index + 1, f"第 {index + 1} 条：请问快递多久能到"])


def _long_text(length: int, seed: int) -> str:
    rng = random.Random(seed)
    sentences = [
        "用户反映家里宽带晚上经常掉线",
        "客服核实后发现光猫型号较旧",
        "建议上门更换设备并检查线路",
        "用户同意周六上午安排工程师",
        "工程师到场后测试信号正常",
        "最后为用户更换了新的光猫",
        "用户表示问题解决并给出好评",
        "回访时用户提到网速比之前快了很多",
    ]
    parts: list[str] = []
    while sum(len(p) for p in parts) < length:
        parts.append(rng.choice(sentences) + rng.choice(["。", "，", "；", "！"]))
    return "".join(parts)[:length]


def make_long_texts(path: Path) -> None:
    lengths = [3000, 6000, 12000, 18000, 25000]
    path.write_text("\n".join(_long_text(n, n) for n in lengths) + "\n", encoding="utf-8")


def make_broken_xlsx(path: Path) -> None:
    path.write_bytes(random.Random(7).randbytes(4096))


def generate(out_dir: Path, *, big: bool = True) -> dict[str, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    files = {
        "intent_200.xlsx": out_dir / "intent_200.xlsx",
        "gbk_1000.csv": out_dir / "gbk_1000.csv",
        "semicolon.csv": out_dir / "semicolon.csv",
        "utf16.txt": out_dir / "utf16.txt",
        "bom.txt": out_dir / "bom.txt",
        "legacy.xls": out_dir / "legacy.xls",
        "long_texts.txt": out_dir / "long_texts.txt",
        "broken.xlsx": out_dir / "broken.xlsx",
    }
    make_intent_xlsx(files["intent_200.xlsx"])
    make_gbk_csv(files["gbk_1000.csv"])
    make_semicolon_csv(files["semicolon.csv"])
    make_text_files(files["utf16.txt"], files["bom.txt"])
    make_legacy_xls(files["legacy.xls"])
    make_long_texts(files["long_texts.txt"])
    make_broken_xlsx(files["broken.xlsx"])
    if big:
        files["big_120k.csv"] = out_dir / "big_120k.csv"
        make_big_csv(files["big_120k.csv"])
    return files


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_OUT
    for name, file in generate(target).items():
        print(f"{name:18s} {file.stat().st_size:>12,d} 字节  {file}")

"""生成测试数据到项目根目录的 testdata/（08 第 10.2 节；02 第 11 节）。

后端测试（backend/tests/conftest.py）和端到端测试（frontend/playwright.config.ts 的 globalSetup）在开始前
发现缺少文件就先运行它。它放在根目录的 scripts/ 而不在 backend/ 或 frontend/ 里，是因为前后端的测试都要用，
谁也不该引用对方目录里的文件（AGENTS.md 第 9 条的明文例外）。testdata/ 不提交。

testdata/.make_testdata.sha256 记录生成时本脚本的 SHA-256。脚本改过（或文件是旧版脚本生成的）时哈希对不上，
下次运行会全部重新生成，不会继续用旧数据。

用后端虚拟环境的 Python 运行（需要 openpyxl、xlwt，其余只用标准库）：

    backend\\.venv\\Scripts\\python.exe scripts\\make_testdata.py              生成缺少的文件，已存在的跳过（脚本改过时全部重新生成）
    backend\\.venv\\Scripts\\python.exe scripts\\make_testdata.py --force      全部重新生成
    backend\\.venv\\Scripts\\python.exe scripts\\make_testdata.py --out DIR    改输出目录

pytest 可以直接 import 本文件，调用 generate(out_dir, force=False)，返回 {文件名: 路径}。
每个文件先写到同目录的临时文件，写完再改名，中途被打断不会留下半个文件被当成“已存在”。

| 文件 | 内容 |
|---|---|
| utf8_big.csv | UTF-8（无 BOM）中文，约 3 MB。列：编号、内容。第 1,048,576 字节正好是一个三字节汉字的中间字节，
|              | 即 1 MB 样本的末尾切在汉字中间（02 第 6.2 节），生成后自行断言 |
| export_edge.csv | UTF-8，列：编号、留言、渠道，10 行。留言依次为：=SUM(A1)、+1 开头、纯数字 -2、-2 开头、@x 开头、
|                 | 含 \\x0b、含 \\x1f、4 万字长文本、=A1+B1 开头、普通文本（04 第 9 节的单元格清理） |
| intent_200.xlsx | 工作表“对话”201 行（表头 + 200 行）：会话ID、用户问题、时间、渠道；用户问题覆盖咨询、办理、投诉、
|                 | 建议、其他 5 类，约 10% 重复，3 行为空。另有工作表“说明”5 行 |
| gbk_1000.csv | GB18030 编码，逗号分隔，1,000 行：编号、客户留言、城市 |
| semicolon.csv | 分号分隔，包含引号内带换行的字段：工单号、问题描述、优先级 |
| utf16.txt、bom.txt | UTF-16 和 UTF-8 BOM 编码的按行文本，第 2 行为空 |
| legacy.xls | 用 xlwt 生成的旧版 Excel，工作表“旧版”：编号、内容、金额，10 行 |
| big_120k.csv | 12 万行（id、text），测试截断和性能 |
| long_texts.txt | 5 行，每行 3,000 到 25,000 字，测试摘要分段 |
| broken.xlsx | 4 KB 随机字节，测试解析失败 |
"""

from __future__ import annotations

import argparse
import codecs
import csv
import hashlib
import io
import os
import random
import sys
from collections.abc import Callable
from datetime import datetime, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_OUT = ROOT / "testdata"
STAMP_NAME = ".make_testdata.sha256"     # 记录生成这批文件的脚本哈希

SAMPLE_BYTES = 1 << 20          # 编码识别取的样本大小：1 MB（02 第 6.2 节）
UTF8_BIG_TARGET = 3 * SAMPLE_BYTES

INTENT_SAMPLES = {
    "咨询": [
        "请问流量套餐怎么收费",
        "会员积分什么时候过期",
        "退货运费谁来承担",
        "发票可以开电子版吗",
        "周末客服几点下班",
        "国际漫游需要提前开通吗",
        "宽带安装要多久",
    ],
    "办理": [
        "帮我开通国际漫游",
        "我要退订每月的视频会员",
        "把套餐改成 99 元的那档",
        "申请修改收货地址",
        "我想注销这个账号",
        "帮我补办一张电话卡",
    ],
    "投诉": [
        "快递三天没动了，客服也不回消息",
        "你们乱扣我的话费，我要投诉",
        "收到的商品是坏的，要求赔偿",
        "装宽带的师傅迟到两个小时，态度还很差",
        "退款申请了一周还没到账",
    ],
    "建议": [
        "建议 App 增加夜间模式",
        "希望能支持按天购买流量包",
        "客服排队太久，建议加一个回拨功能",
        "账单页面能不能显示每天的用量",
    ],
    "其他": ["你好", "嗯嗯", "谢谢", "123456", "哈哈哈"],
}

CHANNELS = ["App", "电话", "网页", "微信", "门店"]

LONG_SENTENCES = [
    "用户反映家里宽带晚上经常掉线",
    "客服核实后发现光猫型号较旧",
    "建议上门更换设备并检查线路",
    "用户同意周六上午安排工程师",
    "工程师到场后测试信号正常",
    "最后为用户更换了新的光猫",
    "用户表示问题解决并给出好评",
    "回访时用户提到网速比之前快了很多",
]

BIG_SENTENCES = [
    "客户来电询问本月账单为什么比上个月多出二十元",
    "客服查询后发现客户开通了一个视频会员的连续包月",
    "客户表示没有印象开通过并要求退订和退款",
    "客服解释了退订流程并提交了退款申请",
    "退款预计在三到五个工作日内原路退回",
    "客户对处理结果表示满意并结束了通话",
    "客户希望以后扣费前能收到短信提醒",
    "客服记录了这条建议并转交给产品部门",
]


# ---------------------------------------------------------------- 各个文件


def _questions(count: int, rng: random.Random) -> list[str]:
    """覆盖 5 类意图，约 10% 重复，3 行为空。"""
    pool = [(intent, text) for intent, texts in INTENT_SAMPLES.items() for text in texts]
    questions: list[str] = []
    for index in range(count):
        _intent, text = pool[index % len(pool)]
        questions.append(f"{text}（第 {index + 1} 条）" if index >= len(pool) else text)
    for index in rng.sample(range(10, count), k=max(1, count // 10)):
        questions[index] = questions[index - rng.randint(1, 9)]
    for index in (count // 4, count // 2, count - 5):
        questions[index] = ""
    return questions


def make_intent_xlsx(path: Path, rows: int = 200) -> None:
    from openpyxl import Workbook

    rng = random.Random(20261021)
    wb = Workbook()
    ws = wb.active
    ws.title = "对话"
    ws.append(["会话ID", "用户问题", "时间", "渠道"])
    start = datetime(2026, 9, 15, 9, 0, 0)
    for index, question in enumerate(_questions(rows, rng)):
        ws.append([10023 + index, question, start + timedelta(minutes=7 * index), CHANNELS[index % len(CHANNELS)]])
    notes = wb.create_sheet("说明")
    notes.append(["项目", "内容"])
    notes.append(["来源", "测试数据"])
    notes.append(["行数", rows])
    notes.append(["意图", "咨询、办理、投诉、建议、其他"])
    notes.append(["备注", "约 10% 重复，3 行为空"])
    wb.save(path)


def make_gbk_csv(path: Path, rows: int = 1000) -> None:
    rng = random.Random(1)
    questions = _questions(rows, rng)
    with path.open("w", encoding="gb18030", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(["编号", "客户留言", "城市"])
        for index, question in enumerate(questions):
            writer.writerow([index + 1, question or "", rng.choice(["北京", "上海", "广州", "成都", "武汉"])])


def make_semicolon_csv(path: Path) -> None:
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh, delimiter=";")
        writer.writerow(["工单号", "问题描述", "优先级"])
        writer.writerow(["A001", "第一行\n第二行：引号内带换行", "高"])
        writer.writerow(["A002", "普通的一行；含中文分号", "中"])
        writer.writerow(["A003", "包含;英文分号的\"引号\"字段", "低"])
        writer.writerow(["A004", "", "低"])


_TEXT_LINES = ["第一行：咨询流量", "", "第三行：办理退订", "第四行：投诉扣费"]

# 文本文件一律 newline="" 原样写入：Windows 的文本模式会把 \n 换成 \r\n，写进去的 \r\n 就成了 \r\r\n，
# 按 \r\n、\r、\n 切行时多出空行。这样各平台生成的字节完全一致。


def make_utf16_txt(path: Path) -> None:
    path.write_text("\r\n".join(_TEXT_LINES) + "\r\n", encoding="utf-16", newline="")


def make_bom_txt(path: Path) -> None:
    path.write_text("\n".join(_TEXT_LINES) + "\n", encoding="utf-8-sig", newline="")


def make_legacy_xls(path: Path) -> None:
    import xlwt

    book = xlwt.Workbook(encoding="utf-8")
    sheet = book.add_sheet("旧版")
    for col, title in enumerate(["编号", "内容", "金额"]):
        sheet.write(0, col, title)
    for row in range(1, 11):
        sheet.write(row, 0, row)
        sheet.write(row, 1, f"旧版 Excel 第 {row} 行：请问账单怎么查")
        sheet.write(row, 2, row * 10.5)
    book.save(str(path))


def make_big_csv(path: Path, rows: int = 120_000) -> None:
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(["id", "text"])
        for index in range(rows):
            writer.writerow([index + 1, f"第 {index + 1} 条：请问快递多久能到"])


def _long_text(length: int, seed: int) -> str:
    rng = random.Random(seed)
    parts: list[str] = []
    total = 0
    while total < length:
        part = rng.choice(LONG_SENTENCES) + rng.choice(["。", "，", "；", "！"])
        parts.append(part)
        total += len(part)
    return "".join(parts)[:length]


def make_long_texts(path: Path) -> None:
    lengths = [3000, 6000, 12000, 18000, 25000]
    path.write_text("\n".join(_long_text(n, n) for n in lengths) + "\n", encoding="utf-8", newline="")


def make_broken_xlsx(path: Path) -> None:
    path.write_bytes(random.Random(7).randbytes(4096))


def _csv_bytes(rows: list[list[str]]) -> bytes:
    buffer = io.StringIO(newline="")
    csv.writer(buffer).writerows(rows)
    return buffer.getvalue().encode("utf-8")


def _cut_mid_char(data: bytes) -> bool:
    """1 MB 样本的最后一个字节（第 1,048,576 字节）是不是一个三字节汉字的中间字节。"""
    last = SAMPLE_BYTES - 1
    return (
        len(data) > SAMPLE_BYTES
        and 0xE0 <= data[last - 1] <= 0xEF
        and 0x80 <= data[last] <= 0xBF
        and 0x80 <= data[last + 1] <= 0xBF
    )


def make_utf8_big_csv(path: Path) -> None:
    """约 3 MB 的 UTF-8 中文 csv。第 1 行数据的“内容”末尾补 0–2 个英文句点调整对齐，
    使 1 MB 的截断点正好落在一个汉字中间：直接严格解码样本会失败，必须用增量解码器（02 第 6.2 节）。"""
    rng = random.Random(3)
    rows: list[list[str]] = []
    size = 0
    while size < UTF8_BIG_TARGET:
        number = len(rows) + 2  # 第 1 行留给对齐用的那一行
        text = "".join(rng.choice(BIG_SENTENCES) + rng.choice(["。", "，", "；"]) for _ in range(rng.randint(3, 6)))
        rows.append([str(number), text])
        size += len(text) * 3 + 8
    head = _csv_bytes([["编号", "内容"]])
    rest = _csv_bytes(rows)
    for pad in range(64):
        first = _csv_bytes([["1", "第 1 行：大于 1 MB 的 UTF-8 中文文件" + "." * pad]])
        data = head + first + rest
        if _cut_mid_char(data):
            break
    else:
        raise RuntimeError("utf8_big.csv：没能把 1 MB 截断点对齐到汉字中间")
    # 自检：整份是合法 UTF-8、没有 BOM；样本严格解码失败在“末尾不完整”，增量解码器能解开
    data.decode("utf-8")
    assert not data.startswith(b"\xef\xbb\xbf")
    sample = data[:SAMPLE_BYTES]
    try:
        sample.decode("utf-8")
    except UnicodeDecodeError as exc:
        assert exc.start >= SAMPLE_BYTES - 2, exc
    else:
        raise AssertionError("utf8_big.csv：1 MB 样本没有切在汉字中间")
    codecs.getincrementaldecoder("utf-8")("strict").decode(sample, final=False)
    path.write_bytes(data)


EDGE_TEXTS = [
    "=SUM(A1)",
    "+1 加一元升级套餐可以吗",
    "-2",
    "-2 号线地铁站旁边的营业厅几点关门",
    "@x 你好，我的账号被锁了",
    "账单里有个看不见的字符\x0b请帮我查一下",
    "退款进度\x1f什么时候到账",
    None,  # 4 万字的长文本，见 make_export_edge_csv
    "=A1+B1 这是一条公式样式的留言",
    "普通的一行：请问宽带怎么续费",
]
EDGE_LONG_CHARS = 40_000


def make_export_edge_csv(path: Path) -> None:
    """导出单元格清理的边界数据（04 第 9 节）：公式前缀、控制字符、超过 32,767 字的单元格。"""
    long_text = _long_text(EDGE_LONG_CHARS, 40)
    assert len(long_text) == EDGE_LONG_CHARS
    rows = [["编号", "留言", "渠道"]]
    for index, text in enumerate(EDGE_TEXTS, start=1):
        rows.append([str(index), long_text if text is None else text, CHANNELS[index % len(CHANNELS)]])
    path.write_bytes(_csv_bytes(rows))


# ---------------------------------------------------------------- 入口

MAKERS: dict[str, Callable[[Path], None]] = {
    "utf8_big.csv": make_utf8_big_csv,
    "export_edge.csv": make_export_edge_csv,
    "intent_200.xlsx": make_intent_xlsx,
    "gbk_1000.csv": make_gbk_csv,
    "semicolon.csv": make_semicolon_csv,
    "utf16.txt": make_utf16_txt,
    "bom.txt": make_bom_txt,
    "legacy.xls": make_legacy_xls,
    "big_120k.csv": make_big_csv,
    "long_texts.txt": make_long_texts,
    "broken.xlsx": make_broken_xlsx,
}


def _make_atomic(path: Path, maker: Callable[[Path], None]) -> None:
    # 临时文件保留原扩展名（openpyxl 按扩展名判断格式），写完再改名
    tmp = path.with_name(f".{path.stem}.{os.getpid()}.tmp{path.suffix}")
    try:
        maker(tmp)
        os.replace(tmp, path)
    finally:
        if tmp.exists():
            tmp.unlink()


def script_digest() -> str:
    """本脚本内容的 SHA-256。frontend/e2e/global-setup.ts 用同样的算法判断 testdata/ 是否过期。"""
    return hashlib.sha256(Path(__file__).read_bytes()).hexdigest()


def generate(out_dir: Path, *, force: bool = False) -> dict[str, Path]:
    """生成全部测试文件到 out_dir，返回 {文件名: 路径}。已存在的文件跳过；force=True，或 out_dir 里记录的
    脚本哈希与本脚本对不上（脚本改过、或是旧版脚本生成的）时全部重新生成。"""
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = out_dir / STAMP_NAME
    digest = script_digest()
    if not force:
        recorded = stamp.read_text(encoding="ascii").strip() if stamp.is_file() else ""
        force = recorded != digest
    files: dict[str, Path] = {}
    for name, maker in MAKERS.items():
        path = out_dir / name
        if force or not path.is_file():
            _make_atomic(path, maker)
        files[name] = path
    stamp.write_text(digest, encoding="ascii")
    return files


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="生成测试数据到 testdata/（08 第 10.2 节）")
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT, help="输出目录，默认项目根目录的 testdata/")
    parser.add_argument("--force", action="store_true", help="已存在的文件也重新生成")
    args = parser.parse_args(argv)
    try:
        files = generate(args.out, force=args.force)
    except ImportError as exc:
        print(f"缺少依赖：{exc}。请用 backend\\.venv\\Scripts\\python.exe 运行（需要 openpyxl、xlwt）。")
        return 1
    for name, file in files.items():
        print(f"{name:18s} {file.stat().st_size:>12,d} 字节  {file}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
